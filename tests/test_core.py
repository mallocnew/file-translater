import io
import json
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.catalog import is_layout_model, layout_language_error
from app.docx_export import markdown_to_docx, parse_markdown
from app.documents import DocumentError, validate_document
from app.main import app, store
from app.mt_uni import parse_glossary
from app.pdf_export import markdown_to_pdf
from app.textutil import ThinkFilter
from app.translate import TranslateOptions, _file_messages, build_prompt, clean_base_url, make_client, resolve_api_key


def test_think_filter_removes_split_tags():
    stream = ThinkFilter()
    assert stream.push("你好") == "你好"
    assert stream.push("<thi") == ""
    assert stream.push("nk>先想一下</thi") == ""
    assert stream.push("nk>正文") == "正文"
    assert stream.finish() == ""

    unfinished = ThinkFilter()
    assert unfinished.push("保留<thi") == "保留"
    assert unfinished.finish() == "<thi"


def test_prompt_and_file_messages():
    options = TranslateOptions(
        model="qwen-long",
        target_lang="英语",
        source_lang="自动检测",
        glossary="模型=model",
        instructions="",
    )
    prompt = build_prompt(options)
    assert "英语" in prompt
    assert "原文件" in prompt
    assert "模型=model" in prompt
    messages = _file_messages("file-fe-1", prompt)
    assert messages[1]["content"] == "fileid://file-fe-1"


def test_base_url_and_client_builder(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("DASHSCOPE_BASE_URL", "https://dashscope.aliyuncs.com/compatible-mode/v1/")
    assert clean_base_url() == "https://dashscope.aliyuncs.com/compatible-mode/v1"
    monkeypatch.setenv("DASHSCOPE_BASE_URL", "http://example.com/v1")
    with pytest.raises(ValueError):
        clean_base_url()
    monkeypatch.delenv("DASHSCOPE_API_KEY", raising=False)
    with pytest.raises(ValueError, match="DASHSCOPE_API_KEY"):
        resolve_api_key()
    monkeypatch.setenv("DASHSCOPE_API_KEY", "sk-server")
    assert resolve_api_key() == "sk-server"
    client = make_client("sk-test", "https://dashscope.aliyuncs.com/compatible-mode/v1")
    client.close()


def test_markdown_export_keeps_structure():
    markdown = "\n".join(
        [
            "# 标题",
            "这是第一段。",
            "还是第一段。",
            "- 一项",
            "- **两项**",
            "",
            "| 术语 | 译文 |",
            "| --- | --- |",
            "| 模型 | model |",
            "",
            "> 引用一句",
            "",
            "```",
            "code",
            "```",
        ]
    )
    blocks = parse_markdown(markdown)
    assert [block["type"] for block in blocks] == ["heading", "paragraph", "ul", "table", "quote", "code"]
    document = Document(io.BytesIO(markdown_to_docx(markdown)))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "标题" in texts
    assert "这是第一段。还是第一段。" in texts
    assert document.tables[0].cell(1, 1).text == "model"
    pdf = markdown_to_pdf(markdown)
    assert pdf.startswith(b"%PDF")


def test_validate_keeps_original_bytes(tmp_path: Path):
    pdf = tmp_path / "plain.pdf"
    pdf.write_bytes(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    validate_document(pdf)

    note = tmp_path / "note.docx"
    document = Document()
    document.add_heading("页译测试", level=1)
    document.add_paragraph("这是一段需要保留的中文。")
    document.save(note)
    original = note.read_bytes()
    validate_document(note)
    assert note.read_bytes() == original

    bad = tmp_path / "fake.pdf"
    bad.write_bytes(b"not-a-pdf")
    with pytest.raises(DocumentError):
        validate_document(bad)


def test_prepare_and_translate_routes(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(store, "root", tmp_path)
    pdf = tmp_path / "upload.pdf"
    pdf.write_bytes(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\ntrailer\n")
    seen = {}

    def fake_run(file_path, filename, options, emit, cancel, live):
        seen["model"] = options.model
        seen["suffix"] = file_path.suffix
        seen["filename"] = filename
        assert cancel is not None
        emit({"type": "token", "text": "Hello"})
        emit({"type": "done", "text": "Hello", "input_tokens": 3, "output_tokens": 1, "chars": 5})

    monkeypatch.setattr("app.main.run_translation", fake_run)
    with TestClient(app) as client:
        config = client.get("/api/config")
        assert config.status_code == 200
        payload = config.json()
        assert "api_key" not in payload
        assert "DASHSCOPE" not in config.text
        assert any(item["id"] == "qwen-mt-uni" and item["mode"] == "layout" for item in payload["models"])
        assert client.get("/api/preview/deadbeefdeadbeefdeadbeefdeadbeef").status_code == 404
        assert client.get("/api/result/deadbeefdeadbeefdeadbeefdeadbeef").status_code == 404
        with pdf.open("rb") as handle:
            response = client.post("/api/prepare", files={"file": ("报表.pdf", handle, "application/pdf")})
        assert response.status_code == 200
        payload = response.json()
        token = payload["token"]
        assert "pages" not in payload
        with client.stream(
            "POST",
            f"/api/translate/{token}",
            json={
                "model": "qwen-long",
                "target_lang": "英语",
            },
        ) as stream:
            body = "".join(stream.iter_text())
    assert "event: token" in body
    assert "Hello" in body
    assert seen["model"] == "qwen-long"
    assert seen["suffix"] == ".pdf"
    assert seen["filename"] == "报表.pdf"


def test_layout_language_and_glossary():
    assert is_layout_model("qwen-mt-uni")
    assert not is_layout_model("qwen-long")
    assert layout_language_error("简体中文", "英语") is None
    assert layout_language_error("自动检测", "法语") is None
    assert layout_language_error("日语", "简体中文")
    assert parse_glossary("transformer=变换器\nattention：注意力\n") == [
        {"src": "transformer", "tgt": "变换器"},
        {"src": "attention", "tgt": "注意力"},
    ]


def test_export_routes_and_layout_result(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(store, "root", tmp_path)
    pdf = tmp_path / "upload.pdf"
    pdf.write_bytes(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\ntrailer\n")
    seen = {}

    def fake_layout(file_path, filename, options, emit, cancel, live):
        seen["model"] = options.model
        dest = file_path.parent / "translated.pdf"
        dest.write_bytes(b"%PDF-1.4 translated")
        meta = json.loads((file_path.parent / "meta.json").read_text(encoding="utf-8"))
        meta["result"] = {"file": "translated.pdf", "filename": "report-英语.pdf"}
        (file_path.parent / "meta.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
        emit({"type": "done", "mode": "layout", "kind": "pdf", "filename": "report-英语.pdf", "input_tokens": 9, "output_tokens": 4})

    monkeypatch.setattr("app.main.run_layout_translation", fake_layout)
    with TestClient(app) as client:
        exported = client.post("/api/export/pdf", json={"markdown": "# Hello\nworld", "filename": "out.pdf"})
        assert exported.status_code == 200
        assert exported.content.startswith(b"%PDF")
        with pdf.open("rb") as handle:
            prepared = client.post("/api/prepare", files={"file": ("report.pdf", handle, "application/pdf")})
        token = prepared.json()["token"]
        with client.stream(
            "POST",
            f"/api/translate/{token}",
            json={"model": "qwen-mt-uni", "target_lang": "英语"},
        ) as stream:
            body = "".join(stream.iter_text())
        assert "layout" in body
        result = client.get(f"/api/result/{token}")
        assert result.status_code == 200
        assert result.content.startswith(b"%PDF-1.4 translated")
    assert seen["model"] == "qwen-mt-uni"
