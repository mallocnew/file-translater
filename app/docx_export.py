from __future__ import annotations

import io
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

_HEADING = re.compile(r"^(#{1,6})\s+(\S.*)$")
_UL = re.compile(r"^[-*]\s+(.+)$")
_OL = re.compile(r"^\d+\.\s+(.+)$")
_RULE = re.compile(r"^(-{3,}|\*{3,}|_{3,})$")
_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def markdown_to_docx(markdown: str) -> bytes:
    document = Document()
    _set_normal_font(document)
    for block in parse_markdown(markdown):
        kind = block["type"]
        if kind == "heading":
            paragraph = document.add_paragraph(style=f"Heading {block['level']}")
            _add_inline(paragraph, block["text"])
        elif kind == "paragraph":
            _add_inline(document.add_paragraph(), block["text"])
        elif kind == "quote":
            paragraph = document.add_paragraph()
            _add_inline(paragraph, block["text"])
            for run in paragraph.runs:
                run.italic = True
        elif kind == "ul":
            for item in block["items"]:
                _add_inline(document.add_paragraph(style="List Bullet"), item)
        elif kind == "ol":
            for item in block["items"]:
                _add_inline(document.add_paragraph(style="List Number"), item)
        elif kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block["text"])
            _font(run, mono=True)
        elif kind == "table":
            rows = block["rows"]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r_index, row in enumerate(rows):
                for c_index, cell in enumerate(row):
                    paragraph = table.cell(r_index, c_index).paragraphs[0]
                    _add_inline(paragraph, cell)
                    if r_index == 0:
                        for run in paragraph.runs:
                            run.bold = True
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def parse_markdown(markdown: str) -> list[dict]:
    lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[dict] = []
    index = 0
    total = len(lines)
    while index < total:
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("```"):
            index += 1
            body: list[str] = []
            while index < total and not lines[index].strip().startswith("```"):
                body.append(lines[index])
                index += 1
            if index < total:
                index += 1
            blocks.append({"type": "code", "text": "\n".join(body)})
            continue
        heading = _HEADING.match(stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            blocks.append({"type": "heading", "level": level, "text": heading.group(2).strip()})
            index += 1
            continue
        if _RULE.fullmatch(stripped):
            index += 1
            continue
        if _is_table(stripped):
            rows: list[list[str]] = []
            while index < total and _is_table(lines[index].strip()):
                if not _is_separator(lines[index]):
                    rows.append(_split_row(lines[index]))
                index += 1
            if rows:
                width = max(len(row) for row in rows)
                blocks.append({"type": "table", "rows": [row + [""] * (width - len(row)) for row in rows]})
            continue
        unordered = _UL.match(stripped)
        ordered = _OL.match(stripped)
        if unordered or ordered:
            pattern = _OL if ordered else _UL
            items: list[str] = []
            while index < total:
                match = pattern.match(lines[index].strip())
                if not match:
                    break
                items.append(match.group(1).strip())
                index += 1
            blocks.append({"type": "ol" if ordered else "ul", "items": items})
            continue
        if stripped.startswith(">"):
            quoted: list[str] = []
            while index < total and lines[index].strip().startswith(">"):
                quoted.append(lines[index].strip()[1:].strip())
                index += 1
            blocks.append({"type": "quote", "text": _join_lines(quoted)})
            continue
        paragraph = [stripped]
        index += 1
        while index < total and lines[index].strip() and not _starts_block(lines[index].strip()):
            paragraph.append(lines[index].strip())
            index += 1
        blocks.append({"type": "paragraph", "text": _join_lines(paragraph)})
    return blocks


def _starts_block(stripped: str) -> bool:
    return bool(
        stripped.startswith("```")
        or _HEADING.match(stripped)
        or _RULE.fullmatch(stripped)
        or _is_table(stripped)
        or _UL.match(stripped)
        or _OL.match(stripped)
        or stripped.startswith(">")
    )


def _is_table(stripped: str) -> bool:
    return stripped.startswith("|") and stripped.endswith("|") and "|" in stripped[1:-1]


def _is_separator(line: str) -> bool:
    cells = _split_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _split_row(line: str) -> list[str]:
    raw = line.strip()
    if raw.startswith("|"):
        raw = raw[1:]
    if raw.endswith("|"):
        raw = raw[:-1]
    return [cell.strip() for cell in raw.split("|")]


def _join_lines(lines: list[str]) -> str:
    parts = [line for line in lines if line]
    if not parts:
        return ""
    text = parts[0]
    for line in parts[1:]:
        if text[-1].isascii() and line[0].isascii():
            text += " " + line
        else:
            text += line
    return text


def _add_inline(paragraph, text: str) -> None:
    position = 0
    for match in _INLINE.finditer(text):
        if match.start() > position:
            _plain(paragraph, text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _font(run, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            _font(run, mono=True)
        position = match.end()
    if position < len(text):
        _plain(paragraph, text[position:])
    if not paragraph.runs:
        _plain(paragraph, "")


def _plain(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    _font(run)


def _font(run, size: int | None = None, bold: bool = False, mono: bool = False) -> None:
    run.bold = bold or run.bold
    run.font.size = Pt(size or (10 if mono else 11))
    run.font.name = "Consolas" if mono else "Calibri"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), run.font.name)
    fonts.set(qn("w:hAnsi"), run.font.name)
    fonts.set(qn("w:eastAsia"), "等线")


def _set_normal_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "等线")
